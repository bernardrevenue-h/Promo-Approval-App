import streamlit as st
import pandas as pd
import numpy as np
from io import BytesIO
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Promo Approval Tool", page_icon="📊", layout="wide")
st.title("📊 Promo Approval Tool")
st.markdown("Upload satu file Excel dengan sheet: Promo Input, ME Per Store, Sales Mix.")
st.divider()

# ── Helpers ──────────────────────────────────────────────────────
def read_sheet(file, sheet_name):
    df = pd.read_excel(file, sheet_name=sheet_name, header=1)
    df.columns = df.columns.str.strip()
    if 'Visit_Purpose_Name' in df.columns:
        df = df.rename(columns={'Visit_Purpose_Name': 'Platform'})
    return df

def round_to_900(price):
    base = round(price / 1000) * 1000
    candidate = base - 100
    return candidate if abs(price - candidate) <= abs(price - (candidate + 1000)) else candidate + 1000

def to_excel_download(df, price_changed_mask, summary_df):
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Write summary first
        summary_df.to_excel(writer, index=False, sheet_name="Summary")
        # Write detail
        df.to_excel(writer, index=False, sheet_name="Detail")

        # ── Format Summary sheet ──────────────────────────────────
        ws_sum = writer.sheets["Summary"]
        yellow = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
        header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col in range(1, 3):
            cell = ws_sum.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
        ws_sum.column_dimensions['A'].width = 35
        ws_sum.column_dimensions['B'].width = 20

        # ── Format Detail sheet ───────────────────────────────────
        ws = writer.sheets["Detail"]
        header_fill2 = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
        row_fill_alt = PatternFill(start_color="EAF2FF", end_color="EAF2FF", fill_type="solid")
        yellow_row   = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

        # Header row
        for col in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill2
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for i, changed in enumerate(price_changed_mask):
            for col in range(1, len(df.columns) + 1):
                cell = ws.cell(row=i + 2, column=col)
                cell.alignment = Alignment(horizontal="center")
                if changed:
                    cell.fill = yellow_row
                elif i % 2 == 1:
                    cell.fill = row_fill_alt

        # Auto width
        for col in ws.columns:
            max_len = max((len(str(cell.value)) for cell in col if cell.value), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 30)

    return buffer.getvalue()

def evaluate(new_final_prices, book_prices, qtys, store_sales, me_store_current, me_sku_current, pc_sku_current):
    pc_new    = (book_prices - new_final_prices) / book_prices
    price_chg = (new_final_prices - book_prices * (1 - pc_sku_current)) / (book_prices * (1 - pc_sku_current) + 1e-9)
    qty_new   = np.maximum(qtys * (1 - price_chg * 2.5), 0)
    sm_new    = (qty_new * book_prices) / store_sales
    sm_new    = sm_new / sm_new.sum() if sm_new.sum() > 0 else sm_new
    sm_cur    = (qtys * book_prices) / store_sales
    sm_cur    = sm_cur / sm_cur.sum()
    pc_store_current   = (pc_sku_current * sm_cur).sum()
    pc_store_new       = (pc_new * sm_new).sum()
    predicted_me_store = me_store_current + ((pc_store_current - pc_store_new) * 0.70)
    me_sku_new         = me_sku_current + ((pc_sku_current - pc_new) * 0.70)
    me_sku_gap         = me_sku_new.max() - me_sku_new.min()
    return predicted_me_store, me_sku_new, me_sku_gap, pc_new, sm_new, qty_new, pc_store_current

def solve_new_prices(book_prices, final_prices, qtys, store_sales, me_store_current, me_sku_current, me_target):
    pc_sku_current = (book_prices - final_prices) / book_prices
    best_result, best_score, best_me_diff, best_gap = None, 999, 999, 999

    for delta in np.arange(-0.20, 0.20, 0.0005):
        new_pc    = np.clip(pc_sku_current + delta, 0, 0.6)
        new_final = np.array([round_to_900(p) for p in book_prices * (1 - new_pc)])
        pred_me, me_sku_new, gap, pc_new, sm_new, qty_new, pc_store_cur = evaluate(
            new_final, book_prices, qtys, store_sales, me_store_current, me_sku_current, pc_sku_current
        )
        me_diff     = abs(pred_me - me_target)
        gap_penalty = max(0, gap - 0.03)
        score       = me_diff * 10 + gap_penalty
        if score < best_score:
            best_score, best_me_diff, best_gap = score, me_diff, gap
            best_result = (new_final, pred_me, pc_new, sm_new, qty_new, me_sku_new, pc_store_cur)

    if best_gap > 0.03:
        new_final_base, _, pc_new_base, _, _, me_sku_new_base, _ = best_result
        me_sku_mean = me_sku_new_base.mean()
        for adj_strength in np.arange(0.001, 0.05, 0.001):
            me_deviation  = me_sku_new_base - me_sku_mean
            pc_adjustment = -me_deviation * adj_strength * 10
            new_pc_adj    = np.clip(pc_new_base + pc_adjustment, 0, 0.6)
            new_final_adj = np.array([round_to_900(p) for p in book_prices * (1 - new_pc_adj)])
            pred_me, me_sku_new, gap, pc_new, sm_new, qty_new, pc_store_cur = evaluate(
                new_final_adj, book_prices, qtys, store_sales, me_store_current, me_sku_current, pc_sku_current
            )
            me_diff     = abs(pred_me - me_target)
            gap_penalty = max(0, gap - 0.03)
            score       = me_diff * 10 + gap_penalty
            if score < best_score:
                best_score, best_me_diff, best_gap = score, me_diff, gap
                best_result = (new_final_adj, pred_me, pc_new, sm_new, qty_new, me_sku_new, pc_store_cur)

    return best_result, best_me_diff, best_gap

# ── Upload File ───────────────────────────────────────────────────
st.subheader("📁 Upload File")
uploaded_file = st.file_uploader(
    "Upload file Excel (.xlsx) - sheet: Promo Input, ME Per Store, Sales Mix",
    type=["xlsx"]
)

df_promo = df_me_store = df_sales = None

if uploaded_file:
    try:
        xl = pd.ExcelFile(uploaded_file)
        required_sheets = ["Promo Input", "ME Per Store", "Sales Mix"]
        missing = [s for s in required_sheets if s not in xl.sheet_names]
        if missing:
            st.error(f"Sheet tidak ditemukan: {', '.join(missing)}")
            st.stop()
        else:
            st.success(f"File berhasil dibaca - sheet: {', '.join(required_sheets)}")
            df_promo    = read_sheet(uploaded_file, "Promo Input")
            df_me_store = read_sheet(uploaded_file, "ME Per Store")
            df_sales    = read_sheet(uploaded_file, "Sales Mix")
    except Exception as e:
        st.error(f"Gagal membaca file: {e}")
        st.stop()

st.divider()

# ── Pilih Minggu & Platform ───────────────────────────────────────
selected_week = selected_platform = None

if df_promo is not None:
    st.subheader("📅 Pilih Minggu & Platform")

    df_promo['Monday of Week']    = pd.to_datetime(df_promo['Monday of Week'])
    df_me_store['monday_of_week'] = pd.to_datetime(df_me_store['monday_of_week'])
    df_sales['monday_of_week']    = pd.to_datetime(df_sales['monday_of_week'])

    weeks_promo    = set(df_promo['Monday of Week'].dropna().dt.date.unique())
    weeks_me_store = set(df_me_store['monday_of_week'].dropna().dt.date.unique())
    weeks_sales    = set(df_sales['monday_of_week'].dropna().dt.date.unique())
    common_weeks   = sorted(weeks_promo & weeks_me_store & weeks_sales, reverse=True)

    if not common_weeks:
        st.warning("Tidak ada minggu yang sama di ketiga sheet.")
        st.stop()

    col1, col2 = st.columns(2)
    with col1:
        selected_week = st.selectbox(
            "Pilih Monday of Week:",
            options=common_weeks,
            format_func=lambda x: x.strftime("%d %b %Y")
        )
    with col2:
        available_platforms = sorted(df_promo['Platform'].dropna().unique().tolist())
        selected_platform   = st.selectbox("Pilih Platform:", options=available_platforms)

st.divider()

# ── ME Target ────────────────────────────────────────────────────
st.subheader("🎯 ME BD GP Target")
me_target_input = st.number_input(
    "Masukkan ME BD GP Target (%)", min_value=0.0, max_value=100.0,
    value=28.0, step=0.1, format="%.1f"
)
me_target = me_target_input / 100

st.divider()

# ── Generate ─────────────────────────────────────────────────────
ready = uploaded_file and selected_week is not None and selected_platform is not None

if st.button("🚀 Generate Output", type="primary", use_container_width=True, disabled=not ready):
    with st.spinner("Memproses data..."):
        try:
            df_promo_w    = df_promo[df_promo['Monday of Week'].dt.date == selected_week].copy()
            df_me_store_w = df_me_store[df_me_store['monday_of_week'].dt.date == selected_week].copy()
            df_sales_w    = df_sales[df_sales['monday_of_week'].dt.date == selected_week].copy()

            if df_promo_w.empty:
                st.error("Tidak ada data Promo Input untuk minggu ini.")
                st.stop()

            # ME Per Store - filter by selected platform
            df_me_plat = (
                df_me_store_w[df_me_store_w['visit_purpose_name'] == selected_platform]
                .rename(columns={
                    'visit_purpose_name': 'Platform',
                    'store_brand_owner':  'Store Brand',
                    'store_me_percent':   'ME_Store_Pct',
                    'gross_sales':        'Store_Sales',
                    'net_sales':          'Net_Sales'
                })
            )
            df_me_plat['Price_Cut_BD_GP'] = (
                (df_me_plat['Store_Sales'] - df_me_plat['Net_Sales']) / df_me_plat['Store_Sales']
            )

            # Net price map: menu_code + net_price (SUMIFS key)
            net_price_map = (
                df_promo_w[df_promo_w['Platform'] == selected_platform][['Menu Code Child', 'Net Price']]
                .drop_duplicates()
                .rename(columns={'Menu Code Child': 'menu_code', 'Net Price': 'promo_net_price'})
            )

            # Sales Mix - filter platform + SUMIFS by menu_code + net_price
            df_sales_filtered = df_sales_w[df_sales_w['visit_purpose_name'] == selected_platform].copy()
            df_sales_filtered = df_sales_filtered.merge(net_price_map, on='menu_code', how='inner')
            df_sales_filtered = df_sales_filtered[
                df_sales_filtered['net_price'] == df_sales_filtered['promo_net_price']
            ]

            df_qty = (
                df_sales_filtered
                .groupby(['menu_code', 'visit_purpose_name', 'store_brand_owner', 'promo_net_price'])['qty_total']
                .sum()
                .reset_index()
                .rename(columns={
                    'menu_code':          'Menu Code Child',
                    'visit_purpose_name': 'Platform',
                    'store_brand_owner':  'Store Brand',
                    'promo_net_price':    'Net Price',
                    'qty_total':          'Qty_Raw'
                })
            )

            # Build base
            output = df_promo_w[df_promo_w['Platform'] == selected_platform].copy()
            output = output.merge(
                df_me_plat[['Platform', 'Store Brand', 'ME_Store_Pct', 'Store_Sales', 'Net_Sales', 'Price_Cut_BD_GP']],
                on=['Platform', 'Store Brand'], how='left'
            )
            # Join qty by menu_code + platform + store brand + net price
            output = output.merge(
                df_qty[['Menu Code Child', 'Platform', 'Store Brand', 'Net Price', 'Qty_Raw']],
                on=['Menu Code Child', 'Platform', 'Store Brand', 'Net Price'], how='left'
            )

            # Qty = qty_raw / Divider
            output['Divider'] = pd.to_numeric(output['Divider'], errors='coerce').fillna(1)
            output['Qty']     = output['Qty_Raw'] / output['Divider']

            # Current calculations
            output['ME_SKU']            = pd.to_numeric(output['M/E (%)'], errors='coerce')
            output['Price_Cut_Current'] = (output['Book Price'] - output['Final Price']) / output['Book Price']
            output['Sales_Mix_Current'] = (output['Qty'] * output['Book Price']) / output['Store_Sales']
            total_sm = output['Sales_Mix_Current'].sum()
            output['Sales_Mix_Current'] = output['Sales_Mix_Current'] / total_sm

            book_prices      = output['Book Price'].values.astype(float)
            final_prices     = output['Final Price'].values.astype(float)
            qtys             = output['Qty'].values.astype(float)
            store_sales      = output['Store_Sales'].iloc[0]
            me_store_current = output['ME_Store_Pct'].iloc[0]
            me_sku_current   = output['ME_SKU'].values.astype(float)
            pc_bd_gp         = output['Price_Cut_BD_GP'].iloc[0]

            pc_weighted_current = (output['Price_Cut_Current'].values * output['Sales_Mix_Current'].values).sum()

            # Solve
            best_result, best_me_diff, best_gap = solve_new_prices(
                book_prices, final_prices, qtys, store_sales,
                me_store_current, me_sku_current, me_target
            )
            new_finals, pred_me, pc_new, sm_new, qty_new, me_sku_new, pc_store_cur = best_result

            output['New_Final_Price'] = new_finals
            output['New_Price_Cut']   = pc_new
            output['New_Sales_Mix']   = sm_new
            output['New_ME_SKU']      = me_sku_new

            price_changed   = output['New_Final_Price'].values != output['Final Price'].values
            pc_weighted_new = (pc_new * sm_new).sum()

            # ── Results metrics ───────────────────────────────────
            st.divider()
            st.subheader("📤 Output")

            m1, m2, m3, m4 = st.columns(4)
            m1.metric("ME BD GP Current",   f"{me_store_current*100:.2f}%")
            m2.metric("ME BD GP Target",    f"{me_target*100:.2f}%")
            m3.metric("Predicted ME BD GP", f"{pred_me*100:.2f}%")
            m4.metric("Diff vs Target",     f"{best_me_diff*100:.3f}%")

            if best_me_diff <= 0.002:
                st.success(f"Predicted ME BD GP {pred_me*100:.2f}% - dalam toleransi +/-0.2% dari target")
            else:
                st.warning(f"Predicted ME BD GP {pred_me*100:.2f}% - di luar toleransi +/-0.2% dari target")

            if best_gap <= 0.03:
                st.success(f"Gap ME BD GP antar SKU = {best_gap*100:.2f}% - dalam batas 3%")
            else:
                st.warning(f"Gap ME BD GP antar SKU = {best_gap*100:.2f}% - melebihi batas 3%")

            # ── Summary Table ─────────────────────────────────────
            st.subheader("📊 Summary ME BD GP")
            summary = pd.DataFrame({
                'Metric': [
                    'ME BD GP Current (Store)',
                    'ME BD GP Target',
                    'Predicted ME BD GP',
                    'Diff vs Target',
                    'Price Cut BD GP (Store)',
                    'Price Cut Weighted - Current',
                    'Price Cut Weighted - New',
                    'Gap ME BD GP antar SKU',
                ],
                'Value': [
                    f"{me_store_current*100:.2f}%",
                    f"{me_target*100:.2f}%",
                    f"{pred_me*100:.2f}%",
                    f"{best_me_diff*100:.3f}%",
                    f"{pc_bd_gp*100:.2f}%",
                    f"{pc_weighted_current*100:.2f}%",
                    f"{pc_weighted_new*100:.2f}%",
                    f"{best_gap*100:.2f}%",
                ]
            })

            def style_summary(row):
                if row['Metric'] in ['Predicted ME BD GP', 'ME BD GP Target']:
                    return ['background-color: #D5F5E3; font-weight: bold'] * 2
                if row['Metric'] == 'Diff vs Target':
                    return ['background-color: #FDEBD0'] * 2
                if 'Price Cut' in row['Metric']:
                    return ['background-color: #EBF5FB'] * 2
                return [''] * 2

            st.dataframe(
                summary.style.apply(style_summary, axis=1),
                use_container_width=True,
                hide_index=True
            )

            st.divider()

            # ── Detail Table ──────────────────────────────────────
            st.subheader("📋 Detail per SKU")
            st.caption("🟡 Baris kuning = harga berubah dari current")

            display = pd.DataFrame({
                'Platform':        output['Platform'],
                'Store Brand':     output['Store Brand'],
                'Menu Name':       output['Menu Name'],
                'Platform Price':  output['Platform Price'].apply(lambda x: f"Rp {x:,.0f}"),
                'Book Price':      output['Book Price'].apply(lambda x: f"Rp {x:,.0f}"),
                'Final Price':     output['Final Price'].apply(lambda x: f"Rp {x:,.0f}"),
                'PC Current':      output['Price_Cut_Current'].apply(lambda x: f"{x*100:.2f}%"),
                'ME BD GP (cur)':  output['ME_SKU'].apply(lambda x: f"{x*100:.2f}%" if pd.notna(x) else "-"),
                'Qty (cur)':       output['Qty'].apply(lambda x: f"{x:,.0f}" if pd.notna(x) else "-"),
                'Sales Mix (cur)': output['Sales_Mix_Current'].apply(lambda x: f"{x*100:.2f}%"),
                'New Final Price':  output['New_Final_Price'].apply(lambda x: f"Rp {x:,.0f}"),
                'PC New':          output['New_Price_Cut'].apply(lambda x: f"{x*100:.2f}%"),
                'New Sales Mix':   output['New_Sales_Mix'].apply(lambda x: f"{x*100:.2f}%"),
                'New ME BD GP':    output['New_ME_SKU'].apply(lambda x: f"{x*100:.2f}%" if pd.notna(x) else "-"),
            })

            # Total row
            total_row = pd.DataFrame([{
                'Platform':        '',
                'Store Brand':     '',
                'Menu Name':       '📌 TOTAL',
                'Platform Price':  '',
                'Book Price':      '',
                'Final Price':     '',
                'PC Current':      f"{pc_weighted_current*100:.2f}%",
                'ME BD GP (cur)':  f"{me_store_current*100:.2f}%",
                'Qty (cur)':       f"{output['Qty'].sum():,.0f}",
                'Sales Mix (cur)': '100.00%',
                'New Final Price':  '',
                'PC New':          f"{pc_weighted_new*100:.2f}%",
                'New Sales Mix':   '100.00%',
                'New ME BD GP':    f"{pred_me*100:.2f}%",
            }])
            display_with_total = pd.concat([display, total_row], ignore_index=True)

            def highlight_rows(row):
                idx = row.name
                if idx == len(display):  # total row
                    return ['background-color: #2C3E50; color: white; font-weight: bold'] * len(row)
                if idx < len(price_changed) and price_changed[idx]:
                    return ['background-color: #FFFF00; color: black'] * len(row)
                if idx % 2 == 1:
                    return ['background-color: #F2F3F4'] * len(row)
                return [''] * len(row)

            st.dataframe(
                display_with_total.style.apply(highlight_rows, axis=1),
                use_container_width=True,
                hide_index=True
            )

            st.download_button(
                label="⬇️ Download Output (Excel)",
                data=to_excel_download(display_with_total, list(price_changed) + [False], summary),
                file_name="promo_approval_output.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

        except Exception as e:
            st.error(f"Error: {e}")
            st.exception(e)

if not ready:
    st.caption("Upload file, pilih minggu & platform dulu untuk mengaktifkan tombol.")
